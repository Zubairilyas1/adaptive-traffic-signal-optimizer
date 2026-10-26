"""Build reports/report.docx from reports/report.md (lightweight Markdown -> DOCX).

Supports: headings, paragraphs, bullet lists, fenced code, tables and images.
Relative image paths are resolved against reports/.

Usage
-----
    .\\.venv\\Scripts\\python.exe scripts\\build_report_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

from src import config  # noqa: E402

REPORT_MD = config.REPORTS_DIR / "report.md"
OUT = config.REPORTS_DIR / "report.docx"


def _clean(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = text.replace("`", "")
    return text.strip()


def _add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Light Grid Accent 1"
    for r in rows:
        cells = table.add_row().cells
        for i in range(n_cols):
            cells[i].text = _clean(r[i]) if i < len(r) else ""


def build() -> int:
    if not REPORT_MD.exists():
        print(f"ERROR: {REPORT_MD} not found")
        return 1

    lines = REPORT_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph("\n".join(block))
            p.style = doc.styles["Normal"]
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i += 1
            continue

        img = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if img:
            rel = img.group(2)
            img_path = (config.REPORTS_DIR / rel) if not Path(rel).is_absolute() else Path(rel)
            if img_path.exists():
                doc.add_picture(str(img_path), width=Inches(6.0))
                if img.group(1):
                    cap = doc.add_paragraph(_clean(img.group(1)))
                    cap.style = doc.styles["Normal"]
                    for run in cap.runs:
                        run.italic = True
                        run.font.size = Pt(9)
            else:
                doc.add_paragraph(f"[missing figure: {rel}]")
            i += 1
            continue

        if line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = []
            for tl in table_lines:
                if set(tl.replace("|", "").strip()) <= {"-", ":", " "}:
                    continue
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                rows.append(cells)
            _add_table(doc, rows)
            continue

        if line.startswith("# "):
            doc.add_heading(_clean(line[2:]), level=0)
        elif line.startswith("## "):
            doc.add_heading(_clean(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(_clean(line[4:]), level=2)
        elif line.strip().startswith(("- ", "* ")):
            doc.add_paragraph(_clean(line.strip()[2:]), style="List Bullet")
        elif line.strip() == "---":
            doc.add_paragraph("")
        elif line.strip():
            doc.add_paragraph(_clean(line))
        i += 1

    doc.save(OUT)
    print(f"report -> {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(build())
